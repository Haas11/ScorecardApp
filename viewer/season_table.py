"""Season Stats sheet -> styled table, with the workbook's colouring (export_season.py).

  - count stats: the season leader's cell is amber;
  - rate stats: blue -> white -> red, white = team average, ends = lowest and
    highest player value (reversed where lower is better);
  - small-sample players (italic in the workbook) get no heat map.
"""
from __future__ import annotations

import io

import openpyxl
import pandas as pd

LEADER = "#FFC000"
BLUE, WHITE, RED = (0x44, 0x72, 0xC4), (0xFF, 0xFF, 0xFF), (0xC0, 0x39, 0x2B)
TEAM_BG = "#D9D9D9"
COUNT_COLS = ["H", "2B", "3B", "HR", "R", "RBI", "BB", "SB"]
HEAT_COLS = ["AVG", "OBP", "SLG", "OPS", "BABIP", "ISO", "BB%", "wOBA", "RC", "OPS+", "BB/K"]
HEAT_COLS_LOW_BETTER = ["K%", "AB/HR"]

# Header tooltips on the season page: what each stat means, in words.
STAT_HELP = {
    "G": "Games played.",
    "PA": "Plate appearances: every time the player came up to bat, whatever the outcome.",
    "AB": "At-bats: plate appearances that count towards the batting average. "
          "Walks, hit-by-pitches and sacrifices are left out.",
    "H": "Hits: singles, doubles, triples and home runs.",
    "2B": "Doubles: hits where the batter reached second base.",
    "3B": "Triples: hits where the batter reached third base.",
    "HR": "Home runs.",
    "R": "Runs: times the player crossed home plate.",
    "RBI": "Runs batted in: runs that scored because of the player's at-bat.",
    "BB": "Walks (base on balls).",
    "K": "Strikeouts.",
    "SB": "Stolen bases.",
    "AVG": "Batting average: the share of at-bats that ended in a hit.",
    "OBP": "On-base percentage: how often the player reached base by a hit, walk or hit-by-pitch.",
    "SLG": "Slugging: bases gained per at-bat through hits. Rewards extra-base hits.",
    "OPS": "On-base plus slugging: one number for getting on base and hitting for power.",
    "BABIP": "Batting average on balls in play: how often a ball put in play (not a home run) "
             "fell for a hit. Very high or low values often come down to luck.",
    "ISO": "Isolated power: extra bases per at-bat, counting only doubles, triples and home runs.",
    "BB%": "Walk rate: the share of plate appearances that ended in a walk.",
    "K%": "Strikeout rate: the share of plate appearances that ended in a strikeout. Lower is better.",
    "wOBA": "Weighted on-base average: like OBP, but each way of reaching base counts by how much "
            "it is worth, so a home run counts more than a walk.",
    "RC": "Runs created: an estimate of how many runs the player's hitting produced for the team.",
    "OPS+": "OPS compared with the league average: 100 is average, above 100 is better.",
    "AB/HR": "At-bats per home run. Lower is better.",
    "BB/K": "Walks per strikeout: plate discipline. Higher is better.",
}


def load(raw: bytes) -> tuple[pd.DataFrame, list[bool]]:
    """(table with the Team row first, small-sample flag per row)."""
    df = pd.read_excel(io.BytesIO(raw), sheet_name="Season Stats")
    ws = openpyxl.load_workbook(io.BytesIO(raw))["Season Stats"]
    italic = {ws.cell(row=r, column=1).value: bool(ws.cell(row=r, column=1).font.italic)
              for r in range(2, ws.max_row + 1)}
    team = df[df["Name"] == "Team"]
    players = df[df["Name"] != "Team"]
    df = pd.concat([team, players], ignore_index=True)
    return df, [italic.get(n, False) for n in df["Name"]]


def _mix(a: tuple, b: tuple, t: float) -> str:
    t = min(max(t, 0.0), 1.0)
    return "#%02X%02X%02X" % tuple(round(x + (y - x) * t) for x, y in zip(a, b))


def _scale(v: float, lo: float, mid: float, hi: float, invert: bool) -> str:
    low, high = (RED, BLUE) if invert else (BLUE, RED)
    if v <= mid:
        return _mix(low, WHITE, (v - lo) / (mid - lo) if mid > lo else 1.0)
    return _mix(WHITE, high, (v - mid) / (hi - mid) if hi > mid else 0.0)


def styles(df: pd.DataFrame, small: list[bool]) -> pd.DataFrame:
    """CSS per cell (same shape as df)."""
    css = pd.DataFrame("", index=df.index, columns=df.columns)
    is_team = df["Name"] == "Team"
    sig = ~is_team & ~pd.Series(small, index=df.index)
    players = ~is_team

    for col in COUNT_COLS:
        if col in df:
            top = df.loc[players, col].max()
            if pd.notna(top) and top > 0:
                css.loc[players & (df[col] == top), col] = f"background-color: {LEADER}; color: black"

    team_row = df[is_team].iloc[0] if is_team.any() else None
    for col, invert in [(c, False) for c in HEAT_COLS] + [(c, True) for c in HEAT_COLS_LOW_BETTER]:
        if col not in df:
            continue
        vals = pd.to_numeric(df.loc[sig, col], errors="coerce").dropna()
        if vals.empty:
            continue
        lo, hi = vals.min(), vals.max()
        mid = team_row[col] if team_row is not None and pd.notna(team_row[col]) else vals.mean()
        for i, v in vals.items():
            css.at[i, col] = f"background-color: {_scale(v, lo, mid, hi, invert)}; color: black"

    css.loc[is_team, :] = f"background-color: {TEAM_BG}; color: black"
    return css
